import base64
import csv
import io
import json
import struct
import zlib
from pathlib import Path

import pytest
from browserlab import export
from browserlab.rpc import Server

ROOT = Path(__file__).resolve().parents[1]
EXP = "data_static_scaling_v013"


def png_url(w=40, h=20):
    raw = b"".join(b"\x00" + bytes([200, 100, 50]) * w for _ in range(h))
    def ch(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d))
    data = b"\x89PNG\r\n\x1a\n" + ch(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + ch(b"IDAT", zlib.compress(raw)) + ch(b"IEND", b"")
    return "data:image/png;base64," + base64.b64encode(data).decode()


def srv(tmp_path):
    return Server(io.StringIO(), tmp_path, ROOT)


def export_call(s, fmt, path, **kw):
    return s.handle(json.dumps({"id": 1, "method": "export.write", "params": {"id": EXP, "format": fmt, "path": str(path), **kw}}))


def test_all_formats_written_with_provenance(tmp_path):
    s = srv(tmp_path)
    for fmt, name in [("csv_trials", "t.csv"), ("csv_conditions", "c.csv"), ("json", "a.json"), ("xlsx", "r.xlsx"),
                      ("pdf", "r.pdf"), ("png", "c.png")]:
        r = export_call(s, fmt, tmp_path / name, images=[png_url()] if fmt in ("pdf", "png") else [])
        assert "result" in r, r
        assert (tmp_path / name).stat().st_size > 50
    assert (tmp_path / "r.pdf").read_bytes().startswith(b"%PDF")
    assert (tmp_path / "c.png").read_bytes().startswith(export.PNG_MAGIC)
    assert not list(tmp_path.glob("*.tmp"))


def test_json_contents_and_source_hash(tmp_path):
    export_call(srv(tmp_path), "json", tmp_path / "a.json")
    d = json.loads((tmp_path / "a.json").read_text())
    src = ROOT / EXP / "core/processed/trial_summary_core.csv"
    assert d["provenance"]["source_sha256"] == export.sha256_file(src)
    assert d["provenance"]["experiment"]["protocol_key"].startswith("core|")
    assert d["analysis"]["n_trials"] == 48 and d["provenance"]["limitations"]


def test_trials_csv_header_and_roundtrip(tmp_path):
    export_call(srv(tmp_path), "csv_trials", tmp_path / "t.csv")
    lines = (tmp_path / "t.csv").read_text().splitlines()
    assert lines[0].startswith("# BrowserLab") and "protocol=" in lines[0]
    rows = list(csv.DictReader(lines[1:]))
    assert len(rows) == 48


def test_conditions_csv_units_in_headers(tmp_path):
    export_call(srv(tmp_path), "csv_conditions", tmp_path / "c.csv")
    header = (tmp_path / "c.csv").read_text().splitlines()[1]
    assert "mean_MB" in header and "ci_lo_MB" in header


def test_xlsx_sheets(tmp_path):
    from openpyxl import load_workbook
    export_call(srv(tmp_path), "xlsx", tmp_path / "r.xlsx")
    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb.sheetnames == ["Provenance", "Trials", "Conditions", "Scaling", "ANOVA"]
    assert wb["Trials"].max_row == 49 and wb["Conditions"].max_row == 17


def test_path_safety(tmp_path):
    s = srv(tmp_path)
    before = (ROOT / EXP / "core/processed/trial_summary_core.csv").read_bytes()
    bad = [ROOT / EXP / "core" / "x.csv", ROOT / EXP / "evil.csv", ROOT / "analysis" / "x.csv", "relative.csv", tmp_path / "nodir" / "x.csv"]
    for b in bad:
        r = export_call(s, "csv_trials", b)
        assert r["error"]["code"] == "invalid_params", b
    assert export_call(s, "csv_trials", tmp_path / "x.txt")["error"]["code"] == "invalid_params"
    assert export_call(s, "exe", tmp_path / "x.exe")["error"]["code"] == "invalid_params"
    assert (ROOT / EXP / "core/processed/trial_summary_core.csv").read_bytes() == before
    assert not (ROOT / EXP / "evil.csv").exists()


def test_image_validation(tmp_path):
    s = srv(tmp_path)
    for bad in ["data:image/png;base64,AAAA", "http://x/y.png", "data:image/png;base64,!!!"]:
        assert export_call(s, "png", tmp_path / "i.png", images=[bad])["error"]["code"] == "invalid_params"
    assert export_call(s, "png", tmp_path / "i.png", images=[])["error"]["code"] == "invalid_params"


def test_existing_file_replaced_atomically(tmp_path):
    d = tmp_path / "a.json"; d.write_text("old")
    export_call(srv(tmp_path), "json", d)
    assert json.loads(d.read_text())["analysis"]


def test_single_rep_pdf_has_no_inference_claims(tmp_path):
    a = {"metric": "m", "metric_label": "L", "n_trials": 4, "conditions": [], "scaling": [], "anova": None, "notes": ["x"],
         "inference": {"allowed": False, "reason": "descriptive only"}}
    info = {"id": "q", "suite": "core", "status": "pilot", "protocol_key": "k", "browsers": [], "workloads": [], "tab_levels": [], "system": {}}
    src = tmp_path / "s.csv"; src.write_text("a\n1\n")
    prov = export.provenance(info, src, a)
    assert export.pdf_bytes(a, prov, []).startswith(b"%PDF")
